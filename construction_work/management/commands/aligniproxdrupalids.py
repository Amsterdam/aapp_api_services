import logging
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from construction_work.models.project_models import Project

logger = logging.getLogger(__name__)

EXCEL_FILE_PATH = Path(__file__).resolve().parents[1] / "projects_dvo_iprox.xlsx"


class Command(BaseCommand):
    """Align iprox and drupal IDs for construction work projects"""

    help = "Align iprox and drupal IDs for construction work projects"

    def handle(self, *args, **kwargs):

        # get all projects from Database
        projects = Project.objects.all()

        # get all ID pairs from the Excel file
        id_pairs = self.read_id_mapping_file()

        for project in projects:
            # find id pair where iprox_id matches the project's iprox_id
            id_pair = next(
                (pair for pair in id_pairs if pair["iprox_id"] == project.foreign_id),
                None,
            )
            if id_pair:
                project.drupal_id = id_pair["drupal_id"]
                project.save()
            else:
                logger.warning(
                    f"No ID pair found for project with Iprox ID: {project.foreign_id}"
                )

    def read_id_mapping_file(self):
        workbook = load_workbook(EXCEL_FILE_PATH, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = next(rows, None)
            if headers is None:
                raise CommandError(f"The Excel file is empty: {EXCEL_FILE_PATH}")

            try:
                iprox_id_column = headers.index("iprox_id")
                drupal_id_column = headers.index("dvo id")
            except ValueError as error:
                raise CommandError(
                    "The Excel file must contain 'iprox_id' and 'dvo id' columns"
                ) from error

            id_pairs = [
                {"iprox_id": row[iprox_id_column], "drupal_id": row[drupal_id_column]}
                for row in rows
            ]
        finally:
            workbook.close()

        logger.info(f"Read {len(id_pairs)} project ID pairs from the Excel file")

        return id_pairs
